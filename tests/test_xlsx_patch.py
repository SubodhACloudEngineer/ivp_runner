import zipfile

import openpyxl
import pytest

from ivp_runner.xlsx_patch import PatchError, patch_cells, read_cell_texts

SHEET = "IVP Test Plan"


@pytest.fixture
def workbook(tmp_path):
    """Small workbook with the features the patcher must respect."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws["B5"] = "Check the WLANs box"
    ws["D5"] = "Not Started"
    ws["A6"] = "=A5+1"
    ws["D6"] = "Not Started"
    ws["F6"] = "keep"
    ws.merge_cells("G7:H8")
    ws["G7"] = "merged"
    ws["A9"] = "row exists"
    wb.create_sheet("Other")["A1"] = "untouched"
    path = tmp_path / "src.xlsx"
    wb.save(path)
    # Add an opaque part, like an embedded Visio object, that must survive byte-for-byte.
    with zipfile.ZipFile(path, "a") as z:
        z.writestr("xl/embeddings/blob.bin", bytes(range(256)) * 10)
    return path


def entries(path):
    with zipfile.ZipFile(path) as z:
        return {i.filename: z.read(i) for i in z.infolist()}


def test_writes_text_and_preserves_every_other_part(workbook, tmp_path):
    dst = tmp_path / "out.xlsx"
    patch_cells(workbook, dst, SHEET, {"D5": "Complete", "H5": "line1\nline2 & <x>"})

    assert read_cell_texts(dst, SHEET, ["D5", "H5", "B5", "F6"]) == {
        "D5": "Complete",
        "H5": "line1\nline2 & <x>",
        "B5": "Check the WLANs box",
        "F6": "keep",
    }
    before, after = entries(workbook), entries(dst)
    assert list(before) == list(after)  # same parts, same order
    changed = [n for n in before if before[n] != after[n]]
    assert changed == ["xl/worksheets/sheet1.xml"]


def test_keeps_cell_style(workbook, tmp_path):
    src = openpyxl.load_workbook(workbook)[SHEET]["D5"].style_id
    dst = tmp_path / "out.xlsx"
    patch_cells(workbook, dst, SHEET, {"D5": "Complete"})
    assert openpyxl.load_workbook(dst)[SHEET]["D5"].style_id == src


def test_inserts_missing_cell_in_column_order(workbook, tmp_path):
    dst = tmp_path / "out.xlsx"
    patch_cells(workbook, dst, SHEET, {"C6": "new", "Z6": "end"})
    ws = openpyxl.load_workbook(dst)[SHEET]
    assert (ws["C6"].value, ws["Z6"].value, ws["F6"].value) == ("new", "end", "keep")
    row = [c.column_letter for c in ws[6] if c.value is not None]
    assert row == sorted(row, key=lambda col: (len(col), col))


def test_source_is_never_written(workbook):
    with pytest.raises(PatchError, match="dst must differ"):
        patch_cells(workbook, workbook, SHEET, {"D5": "x"})


@pytest.mark.parametrize(
    "cells, message",
    [
        ({"A6": "x"}, "holds a formula"),
        ({"H8": "x"}, "inside merged range"),
        ({"D99": "x"}, "row 99 does not exist"),
        ({"d5": "x"}, "invalid cell reference"),
    ],
)
def test_unsafe_edits_are_refused_and_nothing_written(workbook, tmp_path, cells, message):
    dst = tmp_path / "out.xlsx"
    with pytest.raises(PatchError, match=message):
        patch_cells(workbook, dst, SHEET, cells)
    assert not dst.exists()


def test_merged_top_left_is_writable(workbook, tmp_path):
    dst = tmp_path / "out.xlsx"
    patch_cells(workbook, dst, SHEET, {"G7": "ok"})
    assert read_cell_texts(dst, SHEET, ["G7"]) == {"G7": "ok"}


def test_unknown_sheet(workbook, tmp_path):
    with pytest.raises(PatchError, match="not found"):
        patch_cells(workbook, tmp_path / "o.xlsx", "Nope", {"A1": "x"})
