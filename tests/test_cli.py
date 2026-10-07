"""End-to-end `ivp run` against a fake Mist API. No network."""

import json
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest
import yaml

from ivp_runner.cli import main
from ivp_runner.results import TestResult
from ivp_runner.xlsx_patch import read_cell_texts

from .factories import ap, disconnected
from .fixtures_mist import FIXTURES, site_doc, stats_items

ROOT = Path(__file__).resolve().parent.parent
TOKEN = "cli-test-token-0123456789abcdef"
FIXTURE_SITE = yaml.safe_load((FIXTURES / "site_profile.yaml").read_text())["site_id"]
DESCRIPTIONS = {
    57: "Go to Access-points and select the corresponding site",
    58: "Check the WLANs box to confirm the SSIDs that are being broadcasted",
    59: "Check the Power mode:",
    60: "Check the status to confirm the proper configuration is in place",
    61: "Make sure the Ethernet properties are correct, Full duplex and no errors are shown.",
}


class Resp:
    def __init__(self, status, body, headers=None):
        self.status_code = status
        self.content = json.dumps(body).encode() if not isinstance(body, bytes) else body
        self.headers = headers or {}


class FakeMist:
    """Serves AP stats and site details for one site; records calls."""

    def __init__(self, site_id, items, site=None, status=200):
        self.site_id, self.items, self.status = site_id, items, status
        self.site = site or {**site_doc(), "id": site_id}
        self.calls = []

    def get(self, url, *, headers, params, timeout):
        self.calls.append(url)
        if self.status != 200:
            return Resp(self.status, {"detail": "nope"})
        if url.endswith(f"/sites/{self.site_id}"):
            return Resp(200, self.site)
        if url.endswith(f"/sites/{self.site_id}/stats/devices"):
            hdr = {"X-Page-Total": str(len(self.items)), "X-Page-Page": params["page"]}
            return Resp(200, self.items, hdr)
        return Resp(404, {})


@pytest.fixture(autouse=True)
def token(monkeypatch):
    monkeypatch.setenv("MIST_API_TOKEN", TOKEN)


@pytest.fixture
def mop(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "IVP Test Plan"
    for r, text in DESCRIPTIONS.items():
        ws[f"A{r}"] = 1  # numbering in the real MOP repeats; we never use it
        ws[f"B{r}"] = text
        ws[f"D{r}"] = "Not Started"
        ws[f"Z{r}"] = "pad"
    wb.create_sheet("Cover Sheet")["A1"] = "keep me"
    path = tmp_path / "MOP_template.xlsx"
    wb.save(path)
    return path


def profile_file(tmp_path, site_id, **exp):
    base = {
        "uplink_port": "eth0",
        "ssids": [{"ssid": "A", "bands": ["24", "5"]}, {"ssid": "B", "bands": ["5", "6"]}],
        "mgmt_subnet": "192.0.2.0/24",
        "dns_servers": ["192.0.2.53", "198.51.100.53"],
        "min_uplink_speed_mbps": 1000,
    }
    base.update(exp)
    doc = {"schema_version": 1, "site_id": site_id, "site_class": "large", "expectations": base}
    path = tmp_path / f"{site_id}.yaml"
    path.write_text(yaml.safe_dump(doc), encoding="utf-8")
    return path


def argv(site, mop, out, profile=None, *extra):
    a = [
        "run", "--site", site, "--org", "org-1", "--api-host", "api.example",
        "--catalogue", str(ROOT / "catalogue" / "ap.yaml"),
        "--mop", str(mop), "--out", str(out),
        "--mapping", str(ROOT / "catalogue" / "mop_mapping.yaml"),
    ]  # fmt: skip
    if profile:
        a += ["--profile", str(profile)]
    return a + list(extra)


def run_cli(args, fake, capsys):
    code = main(args, _session=fake, _sleep=lambda s: None)
    out = capsys.readouterr()
    return code, out.out, out.err


def only_run_dir(out: Path) -> Path:
    (d,) = [p for p in out.iterdir() if p.is_dir()]
    return d


# ---------------------------------------------------------------- full run on captured data


def test_full_run_on_captured_fixtures(tmp_path, mop, capsys):
    out = tmp_path / "out"
    fake = FakeMist(FIXTURE_SITE, stats_items())
    code, stdout, stderr = run_cli(
        argv(FIXTURE_SITE, mop, out, FIXTURES / "site_profile.yaml"), fake, capsys
    )
    assert code == 1  # captured site has FAILs and no ERROR

    run_dir = only_run_dir(out)
    assert re.fullmatch(r"\d{8}T\d{6}Z-" + FIXTURE_SITE[:8], run_dir.name)
    doc = json.loads((run_dir / "results.json").read_text())
    assert doc["run"]["exit_code"] == 1 and doc["run"]["site_id"] == FIXTURE_SITE
    assert len(doc["results"]) == 6 * 91
    assert all(TestResult.model_validate(r) for r in doc["results"])

    assert (run_dir / "raw" / "manifest.json").is_file()
    assert (run_dir / "raw" / "site_device_stats.start.json").is_file()
    assert (run_dir / "raw" / "site_device_stats.end.json").is_file()
    assert len(list((run_dir / "evidence").rglob("*.png"))) == 6 * 91
    assert "exit code 1" in (run_dir / "run.log").read_text()

    (book,) = run_dir.glob("MOP_*.xlsx")
    cells = read_cell_texts(book, "IVP Test Plan", [f"D{r}" for r in DESCRIPTIONS])
    assert cells == {
        "D57": "Issue Reported",  # one AP disconnected
        "D58": "Issue Reported",  # AP-01 FAILs
        "D59": "Issue Reported",  # one power-constrained AP
        "D60": "Complete",
        "D61": "Issue Reported",  # two APs at 100 Mbps
    }
    assert read_cell_texts(mop, "IVP Test Plan", ["D57"]) == {"D57": "Not Started"}
    with zipfile.ZipFile(book) as z:
        cards = [n for n in z.namelist() if n.startswith("xl/media/ivp_runner_")]
    assert len(cards) == 5  # one row card per mapped row, anchored in H57..H61
    assert "5 evidence pictures" in (run_dir / "run.log").read_text()

    # 1 site request + 1 stats page at start, 1 stats page at end
    assert len(fake.calls) == 3
    assert "TEST ID" in stdout and "AP-05" in stdout and f"workbook:   {book}" in stdout
    assert "TOTAL:" in stdout and "exit code:  1" in stdout
    assert "ivp: " in stderr  # progress goes to stderr


def test_json_mode_emits_only_results_json(tmp_path, mop, capsys):
    prof = profile_file(tmp_path, "site-1")
    items = [ap(1), ap(2, power_constrained=True)]
    code, stdout, stderr = run_cli(
        argv("site-1", mop, tmp_path / "out", prof, "--json"), FakeMist("site-1", items), capsys
    )
    doc = json.loads(stdout)  # nothing but JSON on stdout
    assert code == doc["run"]["exit_code"] == 1
    assert "TEST ID" not in stdout and "ivp: " in stderr
    assert len(doc["results"]) == 12
    assert doc["run"]["counts"] == {"PASS": 11, "FAIL": 1, "SKIP": 0, "ERROR": 0}
    assert doc["run"]["mop"]["output"].endswith(".xlsx")


# ---------------------------------------------------------------- exit codes


def test_exit_0_when_everything_passes(tmp_path, mop, capsys):
    prof = profile_file(tmp_path, "site-1")
    code, stdout, _ = run_cli(
        argv("site-1", mop, tmp_path / "o", prof), FakeMist("site-1", [ap(1)]), capsys
    )
    assert code == 0 and "6 PASS" in stdout


def test_exit_1_on_fail(tmp_path, mop, capsys):
    prof = profile_file(tmp_path, "site-1")
    items = [ap(1, power_constrained=True)]
    assert (
        run_cli(argv("site-1", mop, tmp_path / "o", prof), FakeMist("site-1", items), capsys)[0]
        == 1
    )


def test_exit_2_when_error_even_with_fail(tmp_path, mop, capsys):
    prof = profile_file(tmp_path, "site-1", mgmt_subnet=None)  # AP-03 cannot be judged
    items = [ap(1, power_constrained=True), disconnected(9)]  # and there are FAILs
    code, stdout, _ = run_cli(
        argv("site-1", mop, tmp_path / "o", prof), FakeMist("site-1", items), capsys
    )
    assert code == 2
    assert "expectation_missing" in stdout and "criteria_not_met" in stdout


def test_exit_2_on_api_failure(tmp_path, mop, capsys):
    prof = profile_file(tmp_path, "site-1")
    code, stdout, _ = run_cli(
        argv("site-1", mop, tmp_path / "o", prof), FakeMist("site-1", [], status=401), capsys
    )
    assert code == 2 and "api_error" in stdout


def test_exit_3_missing_token(tmp_path, mop, capsys, monkeypatch):
    monkeypatch.delenv("MIST_API_TOKEN")
    prof = profile_file(tmp_path, "site-1")
    fake = FakeMist("site-1", [ap(1)])
    code, stdout, _ = run_cli(argv("site-1", mop, tmp_path / "o", prof), fake, capsys)
    assert code == 3 and "MIST_API_TOKEN is not set" in stdout
    assert fake.calls == []
    doc = json.loads((only_run_dir(tmp_path / "o") / "results.json").read_text())
    assert doc["run"]["exit_code"] == 3 and "MIST_API_TOKEN" in doc["run"]["tool_error"]


def test_exit_3_missing_profile_points_to_template(tmp_path, mop, capsys):
    code, stdout, _ = run_cli(
        argv("site-1", mop, tmp_path / "o", tmp_path / "nope.yaml"), FakeMist("site-1", []), capsys
    )
    assert code == 3 and "sites/TEMPLATE.yaml" in stdout


def test_exit_3_profile_for_another_site(tmp_path, mop, capsys):
    prof = profile_file(tmp_path, "site-2")
    code, stdout, _ = run_cli(
        argv("site-1", mop, tmp_path / "o", prof), FakeMist("site-1", []), capsys
    )
    assert code == 3 and "is for site site-2" in stdout


def test_exit_3_when_mop_template_differs_but_results_kept(tmp_path, mop, capsys):
    wb = openpyxl.load_workbook(mop)
    wb["IVP Test Plan"]["B59"] = "A different MOP version"
    wb.save(mop)
    prof = profile_file(tmp_path, "site-1")
    code, stdout, _ = run_cli(
        argv("site-1", mop, tmp_path / "o", prof), FakeMist("site-1", [ap(1)]), capsys
    )
    assert code == 3 and "B59 does not start with" in stdout
    run_dir = only_run_dir(tmp_path / "o")
    doc = json.loads((run_dir / "results.json").read_text())
    assert len(doc["results"]) == 6 and doc["run"]["exit_code"] == 3
    assert list(run_dir.glob("MOP_*.xlsx")) == []


def test_exit_3_on_bad_arguments(capsys):
    with pytest.raises(SystemExit) as e:
        main(["run", "--site", "x"])
    assert e.value.code == 3


def test_api_host_is_required(capsys):
    with pytest.raises(SystemExit) as e:
        main(["run", "--site", "s", "--org", "o", "--catalogue", "c", "--mop", "m", "--out", "o"])
    assert e.value.code == 3
    assert "--api-host" in capsys.readouterr().err


# ---------------------------------------------------------------- safety


def test_token_never_written_or_printed(tmp_path, mop, capsys):
    prof = profile_file(tmp_path, "site-1")
    items = [ap(1, echo=TOKEN)]  # even if the API echoes it
    code, stdout, stderr = run_cli(
        argv("site-1", mop, tmp_path / "o", prof), FakeMist("site-1", items), capsys
    )
    assert TOKEN not in stdout + stderr
    for f in (tmp_path / "o").rglob("*"):
        if f.is_file() and f.suffix in {".json", ".log", ".yaml"}:
            assert TOKEN not in f.read_text(encoding="utf-8"), f


def test_no_interactive_prompts_in_source():
    for f in (ROOT / "src").rglob("*.py"):
        assert "input(" not in f.read_text(encoding="utf-8"), f


def test_template_profile_is_valid_once_site_id_filled(tmp_path):
    from ivp_runner.site_profile import SiteProfile

    doc = yaml.safe_load((ROOT / "sites" / "TEMPLATE.yaml").read_text())
    doc["site_id"] = "abc"
    p = SiteProfile.model_validate(doc)
    assert p.expectations.mgmt_subnet is None and p.expectations.uplink_port == "eth0"
