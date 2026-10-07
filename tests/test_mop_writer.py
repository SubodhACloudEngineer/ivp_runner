import hashlib
import io
import posixpath
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
import pytest
from openpyxl.drawing.image import Image as XLImage
from PIL import Image

from ivp_runner.assert_engine import evaluate
from ivp_runner.mop_writer import load_mapping, render_cells, render_pictures, write_results
from ivp_runner.results import Verdict
from ivp_runner.xlsx_patch import (
    NS_REL,
    NS_XDR,
    PatchError,
    Picture,
    _resolve,
    patch_workbook,
    read_cell_texts,
    sheet_part,
)

from .factories import ap, catalogue, context, disconnected, payloads, profile

MAPPING = Path(__file__).resolve().parent.parent / "catalogue" / "mop_mapping.yaml"
SHEET = "IVP Test Plan"
DESCRIPTIONS = {
    57: "Go to Access-points and select the corresponding site",
    58: "Check the WLANs box to confirm the SSIDs",
    59: "Check the Power mode:",
    60: "Check the status to confirm the proper configuration",
    61: "Make sure the Ethernet properties are correct",
}
BLOB = bytes(range(256)) * 40  # stands in for an embedded Visio object


def png(color, size=(60, 40)) -> bytes:
    b = io.BytesIO()
    Image.new("RGB", size, color).save(b, "PNG")
    return b.getvalue()


def shots():
    """Stand-ins for portal screenshots, one per mapped row."""
    return [Picture(f"H{r}", png("navy", (1280, 800)), f"portal:r{r}:AP-1") for r in range(57, 62)]


@pytest.fixture
def results():
    """AP 1 healthy; AP 2 power-constrained; AP 9 disconnected; AP-03 has no subnet."""
    start = [ap(1), ap(2, power_constrained=True), disconnected(9)]
    evs = evaluate(catalogue(), profile(mgmt_subnet=None), *payloads(start), context())
    return [e.result for e in evs]


def make_template(path: Path, with_pictures: bool = True) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.column_dimensions["H"].width = 91  # as in the real MOP
    for r, text in DESCRIPTIONS.items():
        ws[f"A{r}"] = 1  # the MOP's own numbering repeats; never used
        ws[f"B{r}"] = text
        ws[f"D{r}"] = "Not Started"
        ws[f"Z{r}"] = "pad"
        ws.row_dimensions[r].height = 200
    other = wb.create_sheet("Cover Sheet")
    other["A1"] = "keep me"
    if with_pictures:
        ws.add_image(XLImage(io.BytesIO(png("red"))), "B57")
        ws.add_image(XLImage(io.BytesIO(png("blue"))), "F58")
        other.add_image(XLImage(io.BytesIO(png("green"))), "C3")
    wb.save(path)
    with zipfile.ZipFile(path, "a") as z:
        z.writestr("xl/embeddings/Microsoft_Visio_Drawing.vsdx", BLOB)
    return path


@pytest.fixture
def template(tmp_path):
    return make_template(tmp_path / "MOP.xlsx")


def pictures_in(path: Path, sheet: str = SHEET) -> list[dict]:
    """Every picture anchored on ``sheet``: name, from (col,row) 0-based, image sha256."""
    with zipfile.ZipFile(path) as z:
        part = sheet_part(z, sheet)
        rels_path = posixpath.join(
            posixpath.dirname(part), "_rels", posixpath.basename(part) + ".rels"
        )
        if rels_path not in z.namelist():
            return []
        rels = ET.fromstring(z.read(rels_path))
        target = next(
            (
                r.get("Target")
                for r in rels.iter(f"{{{NS_REL}}}Relationship")
                if r.get("Type").endswith("/drawing")
            ),
            None,
        )
        if target is None:
            return []
        drawing = _resolve(part, target)
        d_rels_path = posixpath.join(
            posixpath.dirname(drawing), "_rels", posixpath.basename(drawing) + ".rels"
        )
        d_rels = {
            r.get("Id"): _resolve(drawing, r.get("Target"))
            for r in ET.fromstring(z.read(d_rels_path)).iter(f"{{{NS_REL}}}Relationship")
        }
        root = ET.fromstring(z.read(drawing))
        out = []
        for anchor in root:
            frm = anchor.find(f"{{{NS_XDR}}}from")
            nv = anchor.find(f".//{{{NS_XDR}}}cNvPr")
            blip = anchor.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
            rid = blip.get(
                "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
            )
            out.append(
                {
                    "name": nv.get("name"),
                    "id": nv.get("id"),
                    "at": (
                        int(frm.find(f"{{{NS_XDR}}}col").text),
                        int(frm.find(f"{{{NS_XDR}}}row").text),
                    ),
                    "sha": hashlib.sha256(z.read(d_rels[rid])).hexdigest(),
                    "ext": anchor.find(f"{{{NS_XDR}}}ext"),
                }
            )
        return out


def ours(pics):
    return [p for p in pics if p["name"].startswith("ivp-runner:")]


def theirs(pics):
    return [(p["name"], p["at"], p["sha"]) for p in pics if not p["name"].startswith("ivp-runner:")]


# ---------------------------------------------------------------- rendering


def test_shipped_mapping_loads_and_covers_ap_checks():
    m = load_mapping(MAPPING)
    assert [t for r in m.rows for t in r.test_ids] == [f"AP-0{i}" for i in range(6)]
    assert set(m.verdict_to_status) == set(Verdict)
    assert set(m.verdict_to_status.values()) <= set(m.allowed_statuses)
    assert m.evidence_column == "H"


def test_status_per_row_uses_rollup(results):
    cells = render_cells(load_mapping(MAPPING), results, catalogue())
    assert cells == {
        "D57": "Issue Reported",  # AP-00: AP 9 disconnected
        "D58": "Complete",  # AP-01: PASS + SKIP
        "D59": "Issue Reported",  # AP-02: AP 2 constrained
        "D60": "Not Started",  # AP-03: ERROR, no subnet in profile
        "D61": "Complete",  # AP-04 + AP-05 both PASS (+ SKIP)
    }


def test_one_row_card_per_mapped_row(results):
    pics = render_pictures(load_mapping(MAPPING), results, catalogue())
    assert [p.cell for p in pics] == ["H57", "H58", "H59", "H60", "H61"]
    assert pics[-1].name == "AP-04+AP-05:r61"
    for p in pics:
        img = Image.open(io.BytesIO(p.png))
        assert img.width == 620 and len(p.png) < 40 * 1024
    again = render_pictures(load_mapping(MAPPING), results, catalogue())
    assert [p.png for p in pics] == [p.png for p in again]  # deterministic


# ---------------------------------------------------------------- writing


def test_write_puts_status_in_d_and_screenshots_in_h(results, template, tmp_path):
    dst = tmp_path / "out.xlsx"
    report = write_results(load_mapping(MAPPING), results, catalogue(), template, dst, shots())
    assert read_cell_texts(dst, SHEET, list(report.cells)) == report.cells
    assert report.pictures == ["H57", "H58", "H59", "H60", "H61"]

    pics = pictures_in(dst)
    assert sorted(p["at"] for p in ours(pics)) == [(7, r) for r in range(56, 61)]  # H57..H61
    ids = [p["id"] for p in pics]
    assert len(ids) == len(set(ids))
    # sized within the cell (H width ~637 px, row 200 pt ~266 px), EMU = px * 9525
    for p in ours(pics):
        assert int(p["ext"].get("cx")) <= 637 * 9525 and int(p["ext"].get("cy")) <= 266 * 9525
    # openpyxl can parse the result and sees every picture
    wb = openpyxl.load_workbook(dst)
    assert len(wb[SHEET]._images) == 7 and len(wb["Cover Sheet"]._images) == 1


def test_existing_pictures_and_untouched_parts_are_preserved(results, template, tmp_path):
    dst = tmp_path / "out.xlsx"
    write_results(load_mapping(MAPPING), results, catalogue(), template, dst, shots())
    assert theirs(pictures_in(dst)) == theirs(pictures_in(template))
    assert pictures_in(dst, "Cover Sheet") == [] or theirs(
        pictures_in(dst, "Cover Sheet")
    ) == theirs(pictures_in(template, "Cover Sheet"))

    with zipfile.ZipFile(template) as a, zipfile.ZipFile(dst) as b:
        before = {n: a.read(n) for n in a.namelist()}
        after = {n: b.read(n) for n in b.namelist()}
    changed = sorted(n for n in before if before[n] != after.get(n))
    new = sorted(set(after) - set(before))
    with zipfile.ZipFile(template) as a:
        part = sheet_part(a, SHEET)
    allowed = {part, "xl/drawings/drawing1.xml", "xl/drawings/_rels/drawing1.xml.rels"}
    assert set(changed) <= allowed, changed
    assert all(re.fullmatch(r"xl/media/ivp_runner_.*\.png", n) for n in new), new
    assert after["xl/embeddings/Microsoft_Visio_Drawing.vsdx"] == BLOB
    assert list(after)[: len(before)] == list(before)  # original order kept


def test_rerun_on_own_output_replaces_our_pictures(results, template, tmp_path):
    first, second = tmp_path / "1.xlsx", tmp_path / "2.xlsx"
    write_results(load_mapping(MAPPING), results, catalogue(), template, first, shots())
    write_results(load_mapping(MAPPING), results, catalogue(), first, second, shots())
    pics = pictures_in(second)
    assert len(ours(pics)) == 5
    assert theirs(pics) == theirs(pictures_in(template))
    with zipfile.ZipFile(second) as z:
        media = [n for n in z.namelist() if n.startswith("xl/media/ivp_runner_")]
    assert len(media) == 5  # old cards removed, not left orphaned


def test_sheet_without_a_drawing_gets_one(results, tmp_path):
    src = make_template(tmp_path / "plain.xlsx", with_pictures=False)
    dst = tmp_path / "out.xlsx"
    write_results(load_mapping(MAPPING), results, catalogue(), src, dst, shots())
    assert len(ours(pictures_in(dst))) == 5
    assert len(openpyxl.load_workbook(dst)[SHEET]._images) == 5


# ---------------------------------------------------------------- refusals


def test_refuses_foreign_picture_in_evidence_cell(results, tmp_path):
    src = tmp_path / "mop.xlsx"
    make_template(src)
    wb = openpyxl.load_workbook(src)
    wb[SHEET].add_image(XLImage(io.BytesIO(png("purple"))), "H58")
    wb.save(src)
    dst = tmp_path / "o.xlsx"
    with pytest.raises(PatchError, match="H58 already holds a picture"):
        write_results(load_mapping(MAPPING), results, catalogue(), src, dst, shots())
    assert not dst.exists()


def test_refuses_when_template_rows_moved(results, template, tmp_path):
    wb = openpyxl.load_workbook(template)
    wb[SHEET]["B59"] = "Something else entirely"
    wb.save(template)
    with pytest.raises(PatchError, match="B59 does not start with"):
        write_results(
            load_mapping(MAPPING), results, catalogue(), template, tmp_path / "o.xlsx", shots()
        )


def test_refuses_to_overwrite_human_evidence_text(results, template, tmp_path):
    wb = openpyxl.load_workbook(template)
    wb[SHEET]["H60"] = "engineer notes"
    wb.save(template)
    with pytest.raises(PatchError, match="H60 already holds text a human wrote"):
        write_results(
            load_mapping(MAPPING), results, catalogue(), template, tmp_path / "o.xlsx", shots()
        )


def test_non_png_picture_rejected(template, tmp_path):
    with pytest.raises(PatchError, match="not a PNG"):
        patch_workbook(template, tmp_path / "o.xlsx", SHEET, {}, [Picture("H57", b"GIF89a", "x")])


def test_large_picture_is_scaled_to_fit_cell(template, tmp_path):
    dst = tmp_path / "o.xlsx"
    patch_workbook(template, dst, SHEET, {}, [Picture("H57", png("red", (2000, 1000)), "big")])
    (big,) = ours(pictures_in(dst))
    cx, cy = int(big["ext"].get("cx")), int(big["ext"].get("cy"))
    assert cx <= (637 - 8) * 9525 and abs(cx / cy - 2.0) < 0.01  # fits width, keeps 2:1


def test_without_screenshots_column_h_gets_no_picture(results, template, tmp_path):
    dst = tmp_path / "o.xlsx"
    report = write_results(load_mapping(MAPPING), results, catalogue(), template, dst)
    assert report.pictures == [] and ours(pictures_in(dst)) == []


def test_picture_outside_evidence_cells_is_refused(results, template, tmp_path):
    stray = [Picture("C57", png("red"), "x")]
    with pytest.raises(PatchError, match="outside the evidence cells"):
        write_results(
            load_mapping(MAPPING), results, catalogue(), template, tmp_path / "o.xlsx", stray
        )
